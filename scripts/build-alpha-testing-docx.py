from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "docs/Alpha Testing Black Box gojags risk.docx"

NAVY = "17365D"
BLUE = "2B78C5"
PALE_BLUE = "EAF3FB"
PALE_GRAY = "F5F7FA"
MID_GRAY = "D9E0E7"
TEXT_GRAY = RGBColor(82, 89, 98)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=MID_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=10.5, bold=False, color=None, font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def clear_and_write(cell, text, bold=False, size=9.5, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    set_run(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_borders(style):
    p_pr = style.element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run(run, size=8, color=TEXT_GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run("ALPHA TESTING  |  gojags risk")
    set_run(run, size=8, bold=True, color=TEXT_GRAY)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_para.add_run("Dokumen Black Box Testing  |  Halaman ")
    set_run(run, size=8, color=TEXT_GRAY)
    add_page_field(footer_para)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    remove_style_borders(styles["Title"])
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run(run, size={1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True)
    return paragraph


def add_paragraph(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        set_run(run, bold=True)
    run = paragraph.add_run(text)
    set_run(run)
    return paragraph


def add_manual_list(doc, items, numbered=False):
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.24)
        paragraph.paragraph_format.first_line_indent = Inches(-0.24)
        paragraph.paragraph_format.space_after = Pt(2)
        prefix = f"{index}. " if numbered else "- "
        run = paragraph.add_run(prefix + item)
        set_run(run, size=10.3)


def add_table(doc, headers, rows, widths=None, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        clear_and_write(cell, header, bold=True, size=9, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, header_fill)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            clear_and_write(cells[index], value, size=9, align=align)
            if row_index % 2 == 0:
                set_cell_shading(cells[index], PALE_GRAY)
            if widths:
                cells[index].width = Inches(widths[index])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_signature_table(doc):
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["Disusun oleh", "Diperiksa oleh", "Disetujui oleh"]
    for index, label in enumerate(labels):
        clear_and_write(table.rows[0].cells[index], label, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[index], PALE_BLUE)
        clear_and_write(table.rows[1].cells[index], "\n\n\n", size=9)
        clear_and_write(table.rows[2].cells[index], "Nama:\nTanggal:", size=9)
    doc.add_paragraph()


def test_case(case_id, module, title, priority, role, prerequisites, test_data, steps, expected):
    return {
        "id": case_id,
        "module": module,
        "title": title,
        "priority": priority,
        "role": role,
        "prerequisites": prerequisites,
        "test_data": test_data,
        "steps": steps,
        "expected": expected,
    }


CASES = [
    test_case("AT-AUT-001", "Autentikasi", "Login SSO GOJAGS dengan akun aktif", "P0", "Semua pengguna", "Pengguna belum login dan memiliki akun GOJAGS aktif yang terdaftar pada aplikasi.", "Akun GOJAGS aktif.", ["Buka halaman login.", "Klik Login SSO GOJAGS.", "Selesaikan autentikasi dengan kredensial valid.", "Amati halaman tujuan dan identitas pengguna."], "Pengguna masuk ke Dashboard, sesi terbentuk, dan nama serta email akun yang benar ditampilkan."),
    test_case("AT-AUT-002", "Autentikasi", "Menolak autentikasi yang gagal atau dibatalkan", "P0", "Semua pengguna", "Pengguna belum login.", "Kredensial tidak valid atau proses SSO dibatalkan.", ["Mulai login SSO.", "Masukkan kredensial tidak valid atau batalkan proses.", "Kembali ke aplikasi."], "Aplikasi tidak membuat sesi, menampilkan pesan yang dapat dipahami, dan tetap berada pada alur login."),
    test_case("AT-AUT-003", "Autentikasi", "Melindungi halaman aplikasi tanpa sesi", "P0", "Pengguna belum login", "Tidak ada sesi aktif.", "URL langsung /repositori atau /manajemen-risiko/identifikasi.", ["Buka URL halaman terlindungi secara langsung.", "Amati respons aplikasi."], "Pengguna diarahkan ke halaman login dan data aplikasi tidak terlihat."),
    test_case("AT-AUT-004", "Autentikasi", "Logout mengakhiri sesi", "P0", "Semua pengguna", "Pengguna telah login.", "Tidak ada.", ["Buka menu profil atau sidebar.", "Klik Logout dan konfirmasi bila diminta.", "Gunakan tombol kembali browser atau buka URL terlindungi."], "Sesi berakhir, pengguna kembali ke login, dan halaman terlindungi tidak dapat dibuka dari sesi lama."),
    test_case("AT-AUT-005", "Autentikasi", "Identitas dan role pengguna sesuai", "P1", "Admin dan pengguna biasa", "Dua akun dengan role berbeda tersedia.", "Akun Admin dan akun pengguna non-admin.", ["Login sebagai Admin dan catat menu yang terlihat.", "Logout lalu login sebagai pengguna biasa.", "Bandingkan identitas, menu, dan aksi yang tersedia."], "Nama/email sesuai akun dan menu serta aksi mengikuti role/permission masing-masing."),

    test_case("AT-GLB-001", "Tata letak global", "Navigasi desktop dan penanda menu aktif", "P1", "Semua pengguna", "Pengguna telah login pada viewport desktop.", "Viewport minimal 1280 px.", ["Buka setiap menu utama.", "Buka submenu Manajemen Risiko.", "Pilih salah satu submenu."], "Menu aktif ditandai konsisten; ketika submenu aktif, induk Manajemen Risiko juga aktif dan terbuka."),
    test_case("AT-GLB-002", "Tata letak global", "Navigasi hamburger pada perangkat kecil", "P0", "Semua pengguna", "Pengguna telah login.", "Viewport iPhone 16 sekitar 393 x 852.", ["Ubah viewport ke ukuran ponsel.", "Pastikan hamburger berada di kanan header.", "Buka dan tutup menu.", "Pilih menu dari drawer."], "Drawer tampil dalam area layar, dapat di-scroll vertikal, tidak memicu scroll horizontal halaman, dan menutup setelah navigasi."),
    test_case("AT-GLB-003", "Tata letak global", "Tema terang dan gelap", "P1", "Semua pengguna", "Pengguna telah login.", "Tema Light dan Dark.", ["Buka beberapa halaman yang memiliki tabel, modal, form, dan badge.", "Aktifkan tema gelap.", "Kembali ke tema terang."], "Teks, ikon, input, tabel, modal, badge, garis, dan tombol tetap terbaca pada kedua tema."),
    test_case("AT-GLB-004", "Tata letak global", "Periode data mengubah konteks tahun", "P0", "Semua pengguna", "Tersedia data pada sedikitnya dua tahun.", "Contoh tahun 2025 dan 2026.", ["Catat data pada tahun awal.", "Ubah Periode Data ke tahun lain.", "Buka Dashboard dan modul risiko."], "Label periode berubah dan data yang tampil mengikuti rentang tahun yang dipilih tanpa mencampur data di luar periode."),
    test_case("AT-GLB-005", "Tata letak global", "Notifikasi dan profil tidak memotong header mobile", "P1", "Semua pengguna", "Pengguna telah login pada viewport ponsel.", "Viewport 360 px dan 393 px.", ["Amati header saat drawer tertutup.", "Buka drawer.", "Periksa akses tema, periode data, profil, panduan, dan logout."], "Header utama tetap sederhana; kontrol sekunder tersedia secara rapi di drawer dan tidak terpotong."),
    test_case("AT-GLB-006", "Tata letak global", "Panduan pengguna dapat dibuka", "P1", "Semua pengguna", "File panduan telah dikonfigurasi.", "Tombol Lihat Panduan.", ["Buka menu bantuan atau area panduan.", "Klik Lihat Panduan."], "Dokumen panduan terbuka atau terunduh pada tab/perangkat pengguna tanpa respons kosong."),

    test_case("AT-DSH-001", "Dashboard", "Ringkasan Dashboard berhasil dimuat", "P0", "Semua pengguna", "Tersedia data risiko pada periode terpilih.", "Periode dengan data risiko.", ["Buka Dashboard.", "Tunggu seluruh indikator selesai dimuat.", "Bandingkan beberapa angka dengan data sumber."], "Kartu statistik, grafik, matriks, dan daftar prioritas tampil tanpa error dan konsisten dengan data sumber."),
    test_case("AT-DSH-002", "Dashboard", "Dashboard tanpa data", "P1", "Semua pengguna", "Pilih periode yang tidak memiliki data.", "Tahun tanpa data.", ["Ubah periode ke tahun tanpa data.", "Buka atau muat ulang Dashboard."], "Aplikasi menampilkan nilai nol atau empty state yang jelas, bukan error 500 atau komponen rusak."),

    test_case("AT-KON-001", "Penetapan Konteks", "Navigasi seluruh kelompok data", "P0", "Pengguna berizin konteks", "Pengguna memiliki akses Penetapan Konteks.", "Tim Kerja, Kegiatan, Sasaran, Proses Bisnis, Pemangku Kepentingan, Peraturan Perundangan, dan kelompok lain yang tersedia.", ["Buka Penetapan Konteks.", "Geser pemilih kelompok data pada viewport kecil.", "Pilih setiap tab."], "Seluruh tab dapat dijangkau, termasuk dengan scroll horizontal khusus area tab, dan isi sesuai tab terpilih."),
    test_case("AT-KON-002", "Penetapan Konteks", "Menambah data referensi valid", "P0", "Admin atau pengguna berizin create", "Tab referensi telah dibuka.", "Nama unik dan field wajib yang valid.", ["Klik Tambah pada salah satu tab.", "Isi seluruh field wajib.", "Simpan.", "Cari data yang baru dibuat."], "Data tersimpan, notifikasi sukses tampil, dan pilihan baru tersedia pada modul terkait."),
    test_case("AT-KON-003", "Penetapan Konteks", "Validasi field wajib data referensi", "P1", "Admin atau pengguna berizin create", "Form tambah terbuka.", "Field wajib kosong.", ["Biarkan field wajib kosong.", "Klik Simpan."], "Data tidak tersimpan dan pesan validasi muncul pada field yang perlu diperbaiki."),
    test_case("AT-KON-004", "Penetapan Konteks", "Mengubah data referensi", "P1", "Admin atau pengguna berizin update", "Data referensi uji telah tersedia.", "Nama atau atribut baru.", ["Klik Edit pada data uji.", "Ubah nilai.", "Simpan.", "Muat ulang halaman."], "Perubahan tersimpan dan nilai terbaru tampil setelah reload serta pada pilihan modul terkait."),
    test_case("AT-KON-005", "Penetapan Konteks", "Menghapus data referensi", "P1", "Admin atau pengguna berizin delete", "Data uji tidak sedang dipakai atau tersedia skenario data terpakai.", "Satu data bebas relasi dan satu data yang sudah direferensikan.", ["Hapus data bebas relasi dan konfirmasi.", "Coba hapus data yang masih digunakan."], "Data bebas relasi terhapus; data terpakai ditolak secara aman dengan pesan yang menjelaskan sebabnya."),

    test_case("AT-IDN-001", "Identifikasi Risiko", "Menambah risiko dan menyimpan semua", "P0", "Pengguna berizin identifikasi", "Data konteks dan periode telah tersedia.", "Risiko, penyebab, dampak, kategori, sumber, jenis, sasaran/kegiatan, unit kerja, dan tahun.", ["Buka Identifikasi Risiko.", "Tambah baris.", "Isi seluruh data wajib.", "Klik Simpan Semua.", "Muat ulang halaman."], "Risiko tersimpan pada periode yang benar dan seluruh nilai tetap tampil setelah reload."),
    test_case("AT-IDN-002", "Identifikasi Risiko", "Validasi data identifikasi tidak lengkap", "P0", "Pengguna berizin identifikasi", "Tabel identifikasi dapat diedit.", "Risiko tanpa field wajib.", ["Tambah baris.", "Kosongkan satu atau lebih field wajib.", "Klik Simpan Semua."], "Penyimpanan tidak menghasilkan data tidak lengkap dan field/baris bermasalah diberi pesan yang jelas."),
    test_case("AT-IDN-003", "Identifikasi Risiko", "Mengubah dan menghapus risiko", "P1", "Pengguna berizin update/delete", "Tersedia risiko uji.", "Perubahan uraian risiko dan satu risiko yang aman dihapus.", ["Ubah salah satu nilai risiko lalu simpan.", "Muat ulang dan verifikasi perubahan.", "Hapus risiko uji dengan konfirmasi."], "Perubahan tersimpan; penghapusan hanya terjadi setelah konfirmasi dan daftar diperbarui."),
    test_case("AT-IDN-004", "Identifikasi Risiko", "Mengambil referensi dari Bank Risiko", "P1", "Pengguna berizin identifikasi", "Bank Risiko memiliki data yang relevan.", "Kata kunci risiko yang tersedia di bank.", ["Klik Cari dari Bank Risiko.", "Masukkan kata kunci.", "Pilih satu hasil dan tambahkan ke tabel.", "Simpan Semua."], "Data bank disalin ke baris identifikasi, dapat disesuaikan, dan tersimpan sebagai data periode aktif."),
    test_case("AT-IDN-005", "Identifikasi Risiko", "Unduh template dan import Excel valid", "P1", "Pengguna berizin import", "Data konteks tersedia.", "Template resmi yang diisi dengan dua baris valid.", ["Unduh template import untuk periode aktif.", "Isi dua baris sesuai format.", "Import file.", "Periksa ringkasan hasil dan tabel."], "File diproses, jumlah sukses/gagal jelas, dan baris valid masuk ke tabel tanpa duplikasi tak disengaja."),
    test_case("AT-IDN-006", "Identifikasi Risiko", "Import Excel dengan format salah", "P1", "Pengguna berizin import", "Halaman Identifikasi terbuka.", "File tanpa kolom wajib atau tipe file tidak didukung.", ["Pilih file yang formatnya salah.", "Jalankan import."], "Import ditolak dengan penjelasan kolom/format yang salah dan data lama tidak berubah."),

    test_case("AT-ANL-001", "Analisis Risiko", "Memuat risiko hasil identifikasi", "P0", "Pengguna berizin analisis", "Tersedia risiko teridentifikasi pada periode aktif.", "Satu risiko yang belum dianalisis.", ["Buka Analisis Risiko.", "Cari risiko uji."], "Risiko identifikasi tampil satu kali dan siap diberi penilaian."),
    test_case("AT-ANL-002", "Analisis Risiko", "Menilai risiko inheren", "P0", "Pengguna berizin analisis", "Risiko uji tersedia.", "Nilai kemungkinan, dampak, pengendalian yang ada, dan efektivitas sesuai pilihan aplikasi.", ["Isi parameter analisis pada risiko uji.", "Periksa hasil perhitungan level/warna.", "Klik Simpan Semua.", "Muat ulang."], "Nilai dan level risiko inheren dihitung konsisten, tersimpan, dan tetap tampil setelah reload."),
    test_case("AT-ANL-003", "Analisis Risiko", "Validasi nilai analisis", "P1", "Pengguna berizin analisis", "Risiko uji tersedia.", "Nilai kosong atau di luar pilihan yang diperbolehkan.", ["Kosongkan parameter wajib atau masukkan nilai tidak valid bila dimungkinkan.", "Klik Simpan Semua."], "Aplikasi menolak nilai tidak valid dan tidak menyimpan hasil analisis yang inkonsisten."),

    test_case("AT-EVL-001", "Evaluasi Risiko", "Menilai risiko residual", "P0", "Pengguna berizin evaluasi", "Risiko telah memiliki analisis inheren.", "Nilai kemungkinan dan dampak residual yang valid.", ["Buka Evaluasi Risiko.", "Isi kemungkinan dan dampak residual.", "Periksa level residual.", "Simpan dan muat ulang."], "Level residual dihitung dan tersimpan pada risiko yang benar."),
    test_case("AT-EVL-002", "Evaluasi Risiko", "Memilih respons risiko", "P0", "Pengguna berizin evaluasi", "Risiko telah dianalisis.", "Mengurangi, Mengalihkan, Menghindari, dan Menerima Risiko.", ["Pilih setiap opsi respons pada data uji terpisah.", "Isi field tambahan yang muncul.", "Simpan Semua."], "Respons tersimpan sesuai pilihan; field kondisional hanya muncul dan diwajibkan ketika relevan."),
    test_case("AT-EVL-003", "Evaluasi Risiko", "Menolak evaluasi tidak lengkap", "P1", "Pengguna berizin evaluasi", "Risiko uji tersedia.", "Respons atau nilai residual wajib dikosongkan.", ["Isi sebagian data evaluasi.", "Klik Simpan Semua."], "Data tidak lengkap tidak dianggap selesai dan pengguna mendapat petunjuk koreksi."),

    test_case("AT-RTP-001", "Rencana Penanganan", "Menampilkan risiko yang perlu dikurangi", "P0", "Pengguna berizin RTP", "Terdapat evaluasi dengan beberapa respons berbeda.", "Risiko respons Mengurangi dan Menerima.", ["Buka Rencana Penanganan.", "Bandingkan daftar dengan hasil Evaluasi."], "Risiko dengan respons Mengurangi tampil untuk dibuatkan RTP; risiko dengan respons lain tidak muncul secara keliru."),
    test_case("AT-RTP-002", "Rencana Penanganan", "Menyimpan rencana penanganan lengkap", "P0", "Pengguna berizin RTP", "Risiko dengan respons Mengurangi tersedia.", "Rencana tindakan, output, penanggung jawab, target waktu, anggaran atau field lain yang tersedia.", ["Isi data RTP pada risiko uji.", "Klik Simpan Semua.", "Muat ulang halaman."], "RTP tersimpan lengkap pada risiko yang benar dan tetap tampil setelah reload."),
    test_case("AT-RTP-003", "Rencana Penanganan", "Validasi target waktu dan field wajib RTP", "P1", "Pengguna berizin RTP", "Risiko uji tersedia.", "Tanggal tidak logis atau field wajib kosong.", ["Isi RTP dengan data tidak lengkap atau rentang waktu tidak valid.", "Klik Simpan Semua."], "Aplikasi menolak data tidak valid dan menunjukkan bagian yang harus diperbaiki."),

    test_case("AT-MTX-001", "Matriks Risiko", "Menampilkan posisi inheren dan residual", "P0", "Semua pengguna berizin", "Risiko telah dianalisis dan dievaluasi.", "Satu risiko dengan posisi inheren dan residual berbeda.", ["Buka Matriks Risiko.", "Pilih risiko uji.", "Amati posisi dan arah perpindahan."], "Posisi risiko sesuai nilai kemungkinan/dampak; perpindahan inheren ke residual dapat dibedakan dengan jelas."),
    test_case("AT-MTX-002", "Matriks Risiko", "Filter dan pilihan beberapa risiko", "P1", "Semua pengguna berizin", "Beberapa risiko tersedia pada periode aktif.", "Dua atau lebih risiko.", ["Pilih beberapa risiko.", "Ubah periode atau filter yang tersedia.", "Kosongkan pilihan."], "Matriks memperbarui isi sesuai pilihan dan memiliki empty state yang jelas ketika tidak ada risiko dipilih."),

    test_case("AT-PMN-001", "Pemantauan Risiko", "Daftar RTP dan pagination", "P0", "Pengguna berizin pemantauan", "Lebih banyak RTP daripada ukuran halaman.", "Minimal 21 RTP bila ukuran halaman 20.", ["Buka Pemantauan Risiko.", "Klik halaman berikutnya.", "Ubah jumlah data per halaman.", "Kembali ke halaman pertama."], "Isi tabel berubah sesuai halaman, nomor urut benar, dan kontrol pagination tidak tertutup tombol bantuan."),
    test_case("AT-PMN-002", "Pemantauan Risiko", "Menyimpan realisasi pemantauan", "P0", "Pengguna berizin pemantauan", "RTP aktif tersedia.", "Status keterjadian, realisasi waktu, dan realisasi output.", ["Klik Update Realisasi pada RTP.", "Isi data realisasi.", "Klik Simpan Realisasi.", "Buka kembali modal."], "Realisasi tersimpan dan nilai terbaru tampil pada tabel serta modal."),
    test_case("AT-PMN-003", "Pemantauan Risiko", "Menambah dokumen pendukung melalui tautan", "P1", "Pengguna berizin pemantauan", "Modal realisasi terbuka.", "Judul dan URL HTTPS yang valid.", ["Klik Tambah Dokumen.", "Pilih Tautan.", "Isi judul dan URL.", "Simpan Realisasi."], "Tautan tersimpan, dapat dibuka, dan tercatat sebagai bukti pendukung RTP."),
    test_case("AT-PMN-004", "Pemantauan Risiko", "Mengunggah berkas bukti mitigasi", "P0", "Pengguna berizin pemantauan", "Penyimpanan upload lokal/produksi telah dikonfigurasi.", "PDF berukuran valid.", ["Tambah Dokumen pada modal realisasi.", "Pilih Upload Berkas.", "Unggah PDF.", "Simpan Realisasi.", "Buka Repositori Dokumen kategori Bukti Dukung Mitigasi."], "Berkas berhasil diunggah, dapat dibuka, dan otomatis tersedia di repositori sebagai Bukti Dukung Mitigasi pada tahun risiko yang sesuai."),
    test_case("AT-PMN-005", "Pemantauan Risiko", "Menolak dokumen pendukung tidak valid", "P1", "Pengguna berizin pemantauan", "Modal dokumen terbuka.", "Tipe file atau ukuran melebihi batas, atau URL tidak valid.", ["Pilih file/URL tidak valid.", "Coba unggah atau simpan."], "Aplikasi menolak input dengan pesan batas tipe/ukuran/format dan tidak membuat dokumen rusak."),

    test_case("AT-PLP-001", "Pelaporan Risiko", "Rekap proses risiko lengkap", "P0", "Pengguna berizin pelaporan", "Satu risiko telah melalui identifikasi sampai pemantauan.", "Risiko uji end-to-end.", ["Buka Pelaporan Risiko.", "Cari risiko uji.", "Bandingkan setiap kolom dengan data sumber."], "Rekap menampilkan data identifikasi, analisis, evaluasi, RTP, realisasi, dokumen, dan persetujuan secara konsisten."),
    test_case("AT-PLP-002", "Pelaporan Risiko", "Mengubah status persetujuan", "P0", "Pengguna berizin persetujuan", "Baris pelaporan tersedia.", "Draft, Disetujui, Ditolak, serta nama approver bila relevan.", ["Buka aksi Persetujuan.", "Pilih status dan isi approver.", "Simpan.", "Muat ulang halaman."], "Status dan approver tersimpan, tampil jelas di tabel, dan dapat diperbarui sesuai izin."),
    test_case("AT-PLP-003", "Pelaporan Risiko", "Unduh laporan Excel", "P0", "Pengguna berizin pelaporan", "Data tersedia pada periode aktif.", "Periode satu tahun dan rentang tahun.", ["Klik Unduh Excel.", "Buka file hasil unduhan.", "Periksa judul, periode, header, baris, dan dokumen/persetujuan."], "File Excel dapat dibuka, berisi data periode terpilih, dan struktur kolom tidak bergeser atau terpotong."),
    test_case("AT-PLP-004", "Pelaporan Risiko", "Pagination pelaporan", "P1", "Pengguna berizin pelaporan", "Jumlah data melebihi ukuran halaman.", "Data lebih dari 20 baris.", ["Pindah ke halaman berikutnya.", "Ubah data per halaman.", "Periksa nomor urut dan baris terakhir."], "Tabel berganti sesuai halaman, jumlah data benar, dan kontrol dapat digunakan pada desktop maupun mobile."),

    test_case("AT-BNK-001", "Bank Risiko", "Pencarian risiko berdasarkan kata kunci", "P1", "Pengguna berizin Bank Risiko", "Bank memiliki data historis.", "Kata kunci pada risiko, penyebab, atau dampak.", ["Buka Bank Risiko.", "Masukkan kata kunci.", "Atur jumlah hasil bila tersedia.", "Klik Cari."], "Hasil relevan tampil dengan informasi yang cukup untuk dipakai sebagai referensi."),
    test_case("AT-BNK-002", "Bank Risiko", "Pencarian tanpa hasil", "P2", "Pengguna berizin Bank Risiko", "Halaman Bank Risiko terbuka.", "Kata kunci acak yang tidak tersedia.", ["Masukkan kata kunci tanpa kecocokan.", "Klik Cari."], "Empty state menjelaskan bahwa hasil tidak ditemukan dan halaman tetap stabil."),

    test_case("AT-REP-001", "Repositori Dokumen", "Filter kategori dokumen", "P0", "Pengguna berizin repositori", "Setiap kategori memiliki minimal satu dokumen.", "Semua Berkas, Pedoman dan Kebijakan, Bukti Dukung Mitigasi, Laporan dan Risalah.", ["Klik setiap kartu kategori.", "Periksa daftar dokumen dan penanda kategori aktif."], "Daftar hanya memuat kategori yang dipilih; kategori aktif terlihat jelas pada tema terang dan gelap."),
    test_case("AT-REP-002", "Repositori Dokumen", "Unggah dokumen manual", "P0", "Pengguna berizin create", "Penyimpanan upload telah dikonfigurasi.", "Judul, kategori Bukti Dukung Mitigasi, tahun risiko, dan PDF valid.", ["Klik Upload Dokumen.", "Isi judul, kategori, dan tahun risiko.", "Pilih Upload Berkas dan unggah PDF.", "Simpan Berkas."], "Dokumen tersimpan pada kategori dan tahun yang dipilih, nama uploader benar, dan berkas dapat dibuka."),
    test_case("AT-REP-003", "Repositori Dokumen", "Menambahkan dokumen melalui tautan", "P1", "Pengguna berizin create", "Modal upload terbuka.", "Judul, kategori Laporan dan Risalah, tahun, dan URL HTTPS.", ["Pilih metode Tautan.", "Isi data valid.", "Simpan Berkas.", "Klik Buka pada dokumen baru."], "Dokumen tautan tersimpan dan membuka alamat yang benar secara aman."),
    test_case("AT-REP-004", "Repositori Dokumen", "Validasi form dokumen", "P0", "Pengguna berizin create", "Modal upload terbuka.", "Judul kosong, dokumen kosong, URL salah, atau tahun kosong.", ["Coba simpan setiap kombinasi data tidak valid."], "Penyimpanan ditolak dengan pesan spesifik dan tidak membuat record tanpa dokumen."),
    test_case("AT-REP-005", "Repositori Dokumen", "Mengubah metadata dokumen", "P1", "Admin atau uploader berizin update", "Dokumen manual tersedia.", "Judul, kategori, atau tahun baru.", ["Klik Edit Dokumen.", "Ubah metadata.", "Simpan.", "Filter berdasarkan kategori/tahun baru."], "Perubahan tersimpan dan dokumen berpindah filter sesuai metadata terbaru tanpa kehilangan berkas."),
    test_case("AT-REP-006", "Repositori Dokumen", "Menghapus dokumen dengan konfirmasi", "P1", "Admin atau uploader berizin delete", "Dokumen uji tersedia.", "Satu dokumen manual.", ["Klik Hapus.", "Batalkan konfirmasi dan pastikan dokumen masih ada.", "Ulangi lalu konfirmasi hapus."], "Pembatalan tidak mengubah data; konfirmasi menghapus dokumen dan memperbarui daftar."),
    test_case("AT-REP-007", "Repositori Dokumen", "Pencarian cerdas dengan tahun risiko", "P0", "Pengguna berizin repositori", "Dokumen memiliki tahun dan metadata/ringkasan yang dapat dicari.", "Kata kunci relevan dan tahun risiko tertentu.", ["Masukkan kata kunci.", "Pilih tahun risiko.", "Jalankan/tunggu pencarian.", "Pilih salah satu hasil."], "Hasil hanya berasal dari tahun terpilih dan alasan rekomendasi untuk dokumen terpilih tampil jelas."),
    test_case("AT-REP-008", "Repositori Dokumen", "Membuat ringkasan AI dokumen", "P1", "Pengguna berizin repositori", "BAI_API_KEY valid dan dokumen didukung.", "Dokumen PDF atau teks dengan isi yang dapat diekstrak.", ["Klik Ringkasan pada dokumen.", "Tunggu proses selesai.", "Buka kembali ringkasan."], "Ringkasan dari layanan AI tampil terstruktur, relevan dengan isi dokumen, dan tersimpan/dapat diakses kembali tanpa error provider lama."),
    test_case("AT-REP-009", "Repositori Dokumen", "Pagination mengubah isi tabel", "P0", "Pengguna berizin repositori", "Tersedia lebih dari 10 dokumen.", "Sedikitnya 13 dokumen.", ["Catat daftar pada halaman 1.", "Klik Next atau halaman 2.", "Ubah data per halaman.", "Kembali ke halaman 1."], "Daftar benar-benar berubah, nomor urut melanjutkan halaman, total data akurat, dan bubble bantuan tidak menutupi kontrol."),
    test_case("AT-REP-010", "Repositori Dokumen", "Tampilan repositori pada mobile", "P0", "Pengguna berizin repositori", "Viewport ponsel dan dokumen tersedia.", "Viewport 360 x 800 dan 393 x 852.", ["Buka repositori pada viewport ponsel.", "Periksa kategori, pencarian, daftar, pagination, dan panel rekomendasi.", "Geser tabel jika tampilan tabel digunakan."], "Halaman tidak scroll horizontal; hanya area tabel yang dapat digeser, tidak ada duplikasi tabel/kartu, tombol tidak terpotong, dan panel rekomendasi tersusun vertikal."),

    test_case("AT-FAQ-001", "FAQ dan bantuan", "Mencari dan membuka FAQ", "P1", "Semua pengguna", "FAQ telah terisi di database.", "Kata kunci seperti identifikasi, RTP, upload, atau persetujuan.", ["Buka FAQ.", "Cari dengan kata kunci.", "Buka jawaban yang relevan."], "Hasil tersaring dan jawaban penggunaan aplikasi dapat dibaca dengan baik."),
    test_case("AT-FAQ-002", "FAQ dan bantuan", "Admin menambah, mengubah, dan menghapus FAQ", "P1", "Admin", "Admin memiliki akses FAQ.", "Pertanyaan dan jawaban uji dengan urutan tertentu.", ["Klik Tambah FAQ dan simpan data valid.", "Edit jawaban dan urutannya.", "Hapus FAQ uji dengan konfirmasi."], "Tombol kelola terlihat untuk admin dan seluruh operasi tersimpan di database serta tercermin pada daftar."),
    test_case("AT-FAQ-003", "FAQ dan bantuan", "Chat bantuan menjawab dari FAQ", "P1", "Semua pengguna", "FAQ terkait telah tersedia.", "Pertanyaan penggunaan yang memiliki jawaban FAQ.", ["Buka bubble bantuan.", "Ketik pertanyaan penggunaan.", "Kirim dan baca jawaban."], "Chat menampilkan jawaban yang relevan berdasarkan FAQ dan tetap dapat ditutup tanpa menutupi kontrol utama."),

    test_case("AT-NTF-001", "Notifikasi", "Membuka pusat notifikasi dan urgensi", "P1", "Pengguna berizin notifikasi", "Terdapat notifikasi atau risiko mendekati target.", "Notifikasi baru dan notifikasi lama.", ["Klik ikon notifikasi atau menu Pusat Notifikasi.", "Buka satu notifikasi.", "Kembali ke daftar."], "Daftar dimuat, informasi urgensi jelas, dan tautan membawa pengguna ke konteks data yang benar."),
    test_case("AT-NTF-002", "Notifikasi", "Status baca notifikasi", "P2", "Pengguna berizin notifikasi", "Terdapat notifikasi belum dibaca.", "Satu notifikasi baru.", ["Catat indikator belum dibaca.", "Buka notifikasi.", "Muat ulang halaman."], "Status berubah menjadi dibaca dan indikator jumlah diperbarui secara konsisten."),

    test_case("AT-AKS-001", "Manajemen Akses", "Admin melihat menu administrasi", "P0", "Admin", "Login sebagai admin.", "Tidak ada.", ["Buka drawer/sidebar.", "Perluas Manajemen Akses."], "Menu Pengguna, Role Permissions, Pusat Notifikasi, Audit Log, dan pengelolaan FAQ tampil sesuai kebijakan aplikasi."),
    test_case("AT-AKS-002", "Manajemen Akses", "Mengubah role pengguna", "P0", "Admin", "Pengguna uji dan role tujuan tersedia.", "Pengguna non-admin dan role uji.", ["Buka Pengguna.", "Pilih pengguna uji.", "Ubah role dan simpan.", "Login ulang sebagai pengguna uji."], "Role tersimpan dan hak akses baru berlaku setelah sesi diperbarui."),
    test_case("AT-AKS-003", "Manajemen Akses", "Mengubah permission role", "P0", "Admin", "Role uji tersedia.", "Cabut satu permission create/update pada modul uji.", ["Buka Role Permissions.", "Ubah permission role uji dan simpan.", "Login sebagai pengguna role tersebut.", "Buka modul terkait dan coba akses aksi/API."], "Menu/aksi mengikuti permission; permintaan langsung tanpa izin ditolak, bukan hanya tombol yang disembunyikan."),
    test_case("AT-AKS-004", "Manajemen Akses", "Pengguna biasa tidak dapat membuka administrasi", "P0", "Pengguna non-admin", "Login dengan role tanpa permission admin.", "URL /users, /roles, dan /audit-log.", ["Periksa sidebar.", "Buka URL administrasi secara langsung."], "Menu tidak ditampilkan dan akses langsung ditolak dengan status/pesan yang aman."),

    test_case("AT-AUD-001", "Audit Log", "Mencatat aksi perubahan data", "P0", "Admin/Auditor", "Pengguna melakukan create atau update pada data uji.", "Aksi update risiko atau upload dokumen.", ["Lakukan perubahan data dan catat waktu serta pengguna.", "Buka Audit Log.", "Cari entri terkait.", "Buka detail."], "Log memuat pelaku, waktu, aksi, resource, dan detail perubahan yang memadai tanpa mengekspos rahasia."),
    test_case("AT-AUD-002", "Audit Log", "Filter dan pagination audit", "P1", "Admin/Auditor", "Log memiliki banyak entri.", "Filter pengguna, aksi, resource, dan rentang waktu bila tersedia.", ["Terapkan filter satu per satu.", "Gabungkan beberapa filter.", "Pindah halaman."], "Hasil sesuai filter, total/pagination konsisten, dan detail entri tetap dapat dibuka."),

    test_case("AT-NFR-001", "Kualitas lintas halaman", "Responsif pada ponsel, tablet, dan desktop", "P0", "Semua pengguna", "Data representatif tersedia.", "360 x 800, 393 x 852, 768 x 1024, 1280 x 720, dan 1920 x 1080.", ["Buka setiap menu utama pada seluruh viewport.", "Periksa judul, tombol, tab, tabel, modal, pagination, dan drawer.", "Lakukan interaksi utama pada tiap ukuran."], "Tidak ada elemen penting terpotong; halaman tidak scroll horizontal; tabel besar memiliki scroll lokal; seluruh aksi dapat dijangkau."),
    test_case("AT-NFR-002", "Kualitas lintas halaman", "Kontras dan keterbacaan semua tema", "P0", "Semua pengguna", "Seluruh komponen utama dapat dibuka.", "Light dan Dark mode.", ["Periksa teks normal/sekunder, placeholder, badge, border, ikon, grafik, tabel, dan modal pada kedua tema.", "Fokus pada kolom Risiko Terkait dan status berwarna."], "Semua informasi memiliki kontras memadai dan tidak menyatu dengan latar belakang."),
    test_case("AT-NFR-003", "Kualitas lintas halaman", "Penanganan kegagalan server atau jaringan", "P0", "Semua pengguna", "Dapat mensimulasikan offline atau respons API gagal.", "Putus jaringan saat load dan saat simpan.", ["Putus jaringan sebelum membuka halaman.", "Sambungkan lalu isi form.", "Putus jaringan saat menyimpan."], "Aplikasi menampilkan error yang dapat dipahami, tidak blank/crash, tidak menunjukkan data berhasil tersimpan secara palsu, dan memungkinkan pengguna mencoba lagi."),
    test_case("AT-NFR-004", "Kualitas lintas halaman", "Tidak ada error tak tertangani pada alur utama", "P0", "Tester teknis", "DevTools browser dan log server dapat diamati.", "Satu siklus bisnis end-to-end.", ["Jalankan login sampai pelaporan dan repositori.", "Amati console browser, network, dan log server."], "Tidak ada unhandled rejection, error 500, hydration error, atau request berulang yang tidak diperlukan pada alur berhasil."),
    test_case("AT-NFR-005", "Kualitas lintas halaman", "Keamanan input teks dan tautan", "P0", "Tester teknis", "Form teks dan URL tersedia.", "Teks HTML/script sederhana dan URL javascript yang tidak aman.", ["Masukkan payload uji pada judul, FAQ, risiko, dan URL.", "Simpan lalu tampilkan kembali data."], "Input ditampilkan sebagai teks aman atau ditolak; script tidak dijalankan dan tautan berbahaya tidak dapat digunakan."),
]


MODULE_ORDER = [
    "Autentikasi",
    "Tata letak global",
    "Dashboard",
    "Penetapan Konteks",
    "Identifikasi Risiko",
    "Analisis Risiko",
    "Evaluasi Risiko",
    "Rencana Penanganan",
    "Matriks Risiko",
    "Pemantauan Risiko",
    "Pelaporan Risiko",
    "Bank Risiko",
    "Repositori Dokumen",
    "FAQ dan bantuan",
    "Notifikasi",
    "Manajemen Akses",
    "Audit Log",
    "Kualitas lintas halaman",
]


def add_test_case(doc, case):
    heading = add_heading(doc, f"{case['id']}  {case['title']}", 3)
    heading.paragraph_format.space_before = Pt(9)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rows = [
        ("Prioritas dan peran", f"{case['priority']}  |  {case['role']}"),
        ("Tujuan pengujian", case['title']),
        ("Prasyarat", case['prerequisites']),
        ("Data uji", case['test_data']),
        ("Skenario dan langkah", "\n".join(f"{index}. {step}" for index, step in enumerate(case["steps"], start=1))),
        ("Hasil yang diharapkan", case['expected']),
        ("Hasil aktual", "........................................................................................................\n........................................................................................................"),
        ("Status dan bukti", "[ ] Lulus    [ ] Gagal    [ ] Terblokir    [ ] Tidak Diuji\nBukti/screenshot: ____________________    ID defect: ____________________"),
        ("Tester dan tanggal", "Nama: ______________________________    Tanggal: ____________________"),
    ]
    for row_index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        clear_and_write(cells[0], label, bold=True, size=8.8)
        clear_and_write(cells[1], value, size=8.8)
        set_cell_shading(cells[0], PALE_BLUE)
        if row_index % 2 == 1:
            set_cell_shading(cells[1], PALE_GRAY)
        cells[0].width = Inches(1.45)
        cells[1].width = Inches(5.55)
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Dokumen Alpha Testing gojags risk")
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Black Box Testing dan Checklist Skenario Pengujian")
    set_run(run, size=13, color=TEXT_GRAY)

    doc.add_paragraph()
    add_table(
        doc,
        ["Informasi", "Keterangan"],
        [
            ["Sistem", "gojags risk"],
            ["Jenis pengujian", "Alpha testing dengan pendekatan black box"],
            ["Tujuan", "Memastikan fungsi bisnis, validasi, hak akses, responsivitas, dan penanganan error berjalan sesuai kebutuhan pengguna"],
            ["Lingkungan", "Local/Staging yang menggunakan database dan penyimpanan berkas khusus pengujian"],
            ["Versi build", "______________________________"],
            ["Tanggal pelaksanaan", "______________________________"],
            ["Koordinator pengujian", "______________________________"],
        ],
        widths=[1.65, 5.35],
    )

    add_heading(doc, "Persetujuan Dokumen", 1)
    add_signature_table(doc)

    add_heading(doc, "Tujuan dan Ruang Lingkup", 1)
    add_paragraph(
        doc,
        "Dokumen ini digunakan oleh tim internal untuk memeriksa aplikasi gojags risk sebelum pengujian pengguna atau rilis. "
        "Pengujian dilakukan dari sisi perilaku aplikasi: tester memberikan input, menjalankan aksi, lalu membandingkan keluaran aktual dengan hasil yang diharapkan tanpa menilai implementasi kode di dalamnya.",
    )
    add_paragraph(
        doc,
        "Cakupan meliputi siklus manajemen risiko dari penetapan konteks sampai pelaporan, fitur pendukung, kontrol akses, tema, responsivitas, integrasi dokumen, dan kondisi gagal yang berpotensi menghambat pengguna.",
    )

    add_heading(doc, "Petunjuk Pengisian", 1)
    add_manual_list(
        doc,
        [
            "Jalankan kasus uji sesuai urutan prioritas. P0 harus selesai lebih dahulu, diikuti P1 dan P2.",
            "Gunakan data uji khusus. Hindari mengubah atau menghapus data produksi.",
            "Centang satu status pada setiap kasus: Lulus, Gagal, Terblokir, atau Tidak Diuji.",
            "Isi Hasil Aktual walaupun pengujian lulus. Sertakan bukti screenshot atau rekaman bila diperlukan.",
            "Jika gagal, buat ID defect dan catat langkah reproduksi, dampak, browser, perangkat, serta bukti log yang relevan.",
            "Bersihkan data uji setelah seluruh pengujian selesai apabila tidak dibutuhkan untuk regresi.",
        ],
    )

    add_heading(doc, "Definisi Status dan Prioritas", 1)
    add_table(
        doc,
        ["Kode", "Arti"],
        [
            ["Lulus", "Hasil aktual sama dengan hasil yang diharapkan."],
            ["Gagal", "Ada perbedaan fungsi, validasi, tampilan, atau data dibanding hasil yang diharapkan."],
            ["Terblokir", "Kasus belum dapat dijalankan karena dependensi, data, akses, atau lingkungan belum tersedia."],
            ["Tidak Diuji", "Kasus belum dijalankan pada siklus ini."],
            ["P0", "Kritis. Gangguan dapat menghentikan alur utama, merusak data, membuka akses tanpa izin, atau menyebabkan error server."],
            ["P1", "Penting. Fungsi masih dapat dilanjutkan tetapi pengalaman atau kelengkapan proses terganggu."],
            ["P2", "Pelengkap. Dampak rendah dan tidak menghambat alur utama."],
        ],
        widths=[1.1, 5.9],
    )

    add_heading(doc, "Kriteria Pengujian", 1)
    add_heading(doc, "Kriteria Masuk", 2)
    add_manual_list(
        doc,
        [
            "Build aplikasi dapat dijalankan dan URL pengujian tersedia.",
            "Database pengujian, akun, role, permission, layanan email/notifikasi, BAI API, dan penyimpanan berkas telah dikonfigurasi sesuai skenario yang diuji.",
            "Data referensi minimal tersedia untuk membuat satu risiko end-to-end.",
            "Browser dan ukuran viewport pengujian telah ditentukan.",
            "Tim sepakat mengenai tingkat keparahan dan alur pelaporan defect.",
        ],
    )
    add_heading(doc, "Kriteria Keluar", 2)
    add_manual_list(
        doc,
        [
            "Seluruh kasus P0 telah dijalankan dan lulus.",
            "Tidak ada defect kritis atau tinggi yang masih terbuka pada alur utama.",
            "Kasus P1 yang gagal memiliki keputusan perbaikan atau penerimaan risiko yang terdokumentasi.",
            "Hasil aktual, bukti, dan defect telah direkap serta disetujui koordinator pengujian.",
            "Regresi dilakukan pada area yang terdampak setelah perbaikan.",
        ],
    )

    add_heading(doc, "Matriks Cakupan", 1)
    coverage_rows = []
    for module in MODULE_ORDER:
        module_cases = [case for case in CASES if case["module"] == module]
        ids = [case["id"] for case in module_cases]
        coverage_rows.append([module, str(len(module_cases)), f"{ids[0]} s.d. {ids[-1]}"])
    add_table(doc, ["Modul", "Jumlah", "ID kasus uji"], coverage_rows, widths=[2.55, 0.8, 3.65])

    add_heading(doc, "Rekap Hasil Eksekusi", 1)
    add_table(
        doc,
        ["Modul", "Total", "Lulus", "Gagal", "Terblokir", "Tidak Diuji"],
        [[module, str(len([case for case in CASES if case["module"] == module])), "", "", "", ""] for module in MODULE_ORDER] + [["TOTAL", str(len(CASES)), "", "", "", ""]],
        widths=[2.25, 0.75, 0.9, 0.9, 1.1, 1.1],
    )

    add_heading(doc, "Ringkasan Defect", 1)
    add_table(
        doc,
        ["ID", "Kasus uji", "Ringkasan", "Severity", "Status", "PIC"],
        [["", "", "", "", "", ""] for _ in range(5)],
        widths=[0.75, 1.15, 2.6, 0.85, 0.85, 0.8],
    )

    doc.add_page_break()
    add_heading(doc, "Checklist dan Skenario Pengujian", 1)
    add_paragraph(
        doc,
        "Jalankan kasus uji berikut pada lingkungan alpha. Urutan di dalam modul disusun dari alur paling penting ke pemeriksaan pelengkap.",
    )

    for module in MODULE_ORDER:
        add_heading(doc, module, 2)
        for case in [item for item in CASES if item["module"] == module]:
            add_test_case(doc, case)

    doc.add_page_break()
    add_heading(doc, "Catatan Akhir Pengujian", 1)
    add_paragraph(doc, "Ringkasan hasil dan keputusan rilis:")
    for _ in range(8):
        add_paragraph(doc, "................................................................................................................................................")
    add_heading(doc, "Keputusan", 2)
    add_manual_list(doc, ["[ ] Layak dilanjutkan ke UAT", "[ ] Layak dengan perbaikan minor", "[ ] Belum layak dan perlu pengujian ulang"])
    add_signature_table(doc)

    doc.core_properties.title = "Dokumen Alpha Testing gojags risk"
    doc.core_properties.subject = "Black box testing dan checklist skenario pengujian"
    doc.core_properties.author = "Tim gojags risk"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(CASES)} test cases")


if __name__ == "__main__":
    build_document()
